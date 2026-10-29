# -*- coding: utf-8 -*-
"""
Scopus 논문 키프레이즈 빈도 집계 + 워드클라우드 생성기

제목(Title) · 초록(Abstract) · 저자키워드(Author Keywords)에서 키프레이즈를 뽑아
노출빈도를 집계하고, 엑셀 표와 워드클라우드 이미지를 만듭니다.

  - 단수/복수형 통합 : models / model, analyses / analysis, mice / mouse, LLMs / LLM ...
  - 다단어 구문 인식 : "large language model", "quality of life" 등을 하나의 키프레이즈로 집계
      1) 코퍼스 전체의 저자키워드를 구문 사전으로 사용
      2) 제목/초록의 명사구(형용사·명사 연속) 중 여러 논문에 반복 등장하는 구문을 자동 추출
      3) config/custom_phrases.txt 에 적은 구문
    → 제목/초록은 사전에 대해 '가장 긴 구문 우선'으로 매칭하고, 남은 명사는 단어로 집계
  - 논문 내 약어 정의("large language models (LLMs)")를 찾아 약어를 원래 표현으로 치환

사용 예)
  python keyphrase_cloud.py "YRS 2025_scopus.xlsx"
  python keyphrase_cloud.py "YRS 2026_scopus.xlsx" --label YRS2026 --year 2026
"""
from __future__ import annotations

import argparse
import datetime as dt
import hashlib
import pickle
import re
import sys
from collections import Counter, defaultdict
from pathlib import Path

import pandas as pd

BASE_DIR = Path(__file__).resolve().parent
CONFIG_DIR = BASE_DIR / "config"
CACHE_DIR = BASE_DIR / ".cache"
CACHE_VERSION = 1
SPACY_MODEL = "en_core_web_sm"

COL_TITLE, COL_ABSTRACT, COL_KEYWORDS = "Title", "Abstract", "Author Keywords"
SECTION_LABELS = {"title": "제목", "abstract": "초록", "keywords": "저자키워드"}

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(errors="replace")

_logger = None


def set_logger(fn) -> None:
    """진행 메시지를 받을 함수 지정 (웹 화면용). None이면 콘솔 출력."""
    global _logger
    _logger = fn


def log(*parts) -> None:
    msg = " ".join(str(p) for p in parts)
    if _logger:
        _logger(msg)
    else:
        print(msg)


# ─────────────────────────────────────────────────────────────────────────────
# 1. 데이터 읽기
# ─────────────────────────────────────────────────────────────────────────────
def read_table(path: Path, sheet: str | None) -> pd.DataFrame:
    required = (COL_TITLE, COL_ABSTRACT, COL_KEYWORDS)
    if path.suffix.lower() == ".csv":
        df = pd.read_csv(path, dtype=str, encoding="utf-8-sig", keep_default_na=False)
        return df.loc[:, ~df.columns.duplicated()]

    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet and sheet not in wb.sheetnames:
            raise SystemExit(f"시트 '{sheet}'을(를) 찾을 수 없습니다. 시트 목록: {wb.sheetnames}")
        for name in [sheet] if sheet else wb.sheetnames:
            rows = wb[name].iter_rows(values_only=True)
            header = next(rows, None)
            if not header:
                continue
            header = ["" if h is None else str(h).strip() for h in header]
            if not all(c in header for c in required):
                if sheet:
                    raise SystemExit(f"시트 '{name}'에 {required} 열이 없습니다.")
                continue
            width = max(i for i, h in enumerate(header) if h) + 1
            data, blank_run = [], 0
            for row in rows:
                row = list(row[:width]) + [None] * (width - len(row))
                if all(v is None or str(v).strip() == "" for v in row):
                    blank_run += 1
                    if blank_run >= 1000:  # 서식만 남은 빈 행이 계속되면 중단
                        break
                    continue
                blank_run = 0
                data.append(["" if v is None else str(v) for v in row])
            log(f"      시트 '{name}' 사용 ({len(data):,}행)")
            df = pd.DataFrame(data, columns=header[:width])
            return df.loc[:, ~df.columns.duplicated()]
    finally:
        wb.close()
    raise SystemExit(f"{path.name}: {required} 열이 있는 시트를 찾지 못했습니다.")


def load_records(paths: list[str], sheet: str | None, years: list[int] | None):
    frames = []
    for p in map(Path, paths):
        if not p.exists():
            raise SystemExit(f"파일이 없습니다: {p}")
        log(f"[1/5] 데이터 읽는 중: {p.name}")
        frames.append(read_table(p, sheet))
    df = pd.concat(frames, ignore_index=True).fillna("")

    if "Year" in df.columns:
        df["Year"] = df["Year"].astype(str).str.extract(r"(\d{4})", expand=False).fillna("")
        if years:
            df = df[df["Year"].isin([str(y) for y in years])]
        log("      연도 분포:", dict(Counter(df["Year"]).most_common()))
    elif years:
        log("      [주의] Year 열이 없어 연도 필터를 적용하지 못했습니다.")

    before = len(df)
    title_key = df[COL_TITLE].str.lower().str.strip()
    if "EID" in df.columns:
        eid = df["EID"].str.strip()
        dedup_key = eid.where(eid != "", title_key)
    else:
        dedup_key = title_key
    df = df[~dedup_key.duplicated()]
    df = df[(df[COL_TITLE].str.strip() != "") | (df[COL_ABSTRACT].str.strip() != "")]
    return df.reset_index(drop=True), before - len(df)


# ─────────────────────────────────────────────────────────────────────────────
# 2. 텍스트 전처리
# ─────────────────────────────────────────────────────────────────────────────
HTML_RE = re.compile(r"<[^>]+>")
COPYRIGHT_RE = re.compile(r"(?:©|\bCopyright\b\s*(?:©\s*)?\(?\d{4}|\(C\)\s*\d{4}).*$", re.S | re.I)
CITATION_RE = re.compile(r"\((?:[A-Z][A-Za-z.&]*\s){1,6}\d{4};[^)]*\)")
STAT_PAREN_RE = re.compile(
    r"\s*\((?=[^()]*(?:[=<>%±≤≥]|\d\s*[,;–-]\s*\d|\d\.\d|\b(?:CI|OR|HR|IQR|SD|RR|aOR|aHR|vs)\b))[^()]*\)"
)
HEADING_RE = re.compile(r"(?:^|(?<=[.!?;]))\s*([A-Z][A-Za-z /&,\-]{0,60}?)\s*:(?!\d)")
HEADING_WORDS = set("""
background backgrounds aim aims objective objectives purpose purposes introduction method methods
methodology material materials result results conclusion conclusions discussion finding findings
design setting settings participant participants patient patients measure measures main outcome
outcomes intervention interventions exposure exposures evidence summary interpretation significance
importance context rationale trial registration identifier funding limitation limitations
implication implications practical social originality value research study hypothesis and of the
clinical relevance key point points question questions highlights statement problem approach
observation observations case presentation description data source sources selection extraction
synthesis eligibility criteria review level graphical abstract number new what is known this adds
subjects analysis statistical experimental procedures principal major brief scope perspective
perspectives future direction directions novelty impact recommendations benefits
""".split())
EMPTY_ABSTRACT = "[no abstract available]"


def strip_headings(text: str) -> str:
    def repl(m):
        words = re.findall(r"[A-Za-z]+", m.group(1).lower())
        return " " if words and all(w in HEADING_WORDS for w in words) else m.group(0)

    return HEADING_RE.sub(repl, text)


def clean_abstract(text: str) -> str:
    if text.strip().lower() == EMPTY_ABSTRACT:
        return ""
    text = HTML_RE.sub("", text)
    text = COPYRIGHT_RE.sub("", text)
    text = CITATION_RE.sub("", text)
    return strip_headings(text)


# ── 약어 정의 탐지 (Schwartz & Hearst 2003 알고리즘) ──
SHORT_FORM_RE = re.compile(r"\(\s*([A-Za-z][\w\-+/&.]{1,11})\s*\)")
LONG_FORM_STOP = {"the", "a", "an", "of", "and", "in", "on", "for", "to", "with", "by", "as", "or"}


def _long_form(short: str, text: str) -> str | None:
    s, l = len(short) - 1, len(text) - 1
    while s >= 0:
        c = short[s].lower()
        if not c.isalnum():
            s -= 1
            continue
        while l >= 0 and (text[l].lower() != c or (s == 0 and l > 0 and text[l - 1].isalnum())):
            l -= 1
        if l < 0:
            return None
        l -= 1
        s -= 1
    start = text.rfind(" ", 0, l + 1) + 1
    return text[start:].strip()


def find_acronyms(texts) -> dict[str, str]:
    amap: dict[str, str] = {}
    for text in texts:
        for m in SHORT_FORM_RE.finditer(text):
            short = m.group(1).rstrip(".")
            if short.lower() == short or len(short) < 2:
                continue
            before = text[: m.start()]
            before = before[max(before.rfind(ch) for ch in ".;:!?()[],") + 1 :]
            words = before.split()
            cand = " ".join(words[-min(len(short) + 5, len(short) * 2) :])
            long = _long_form(short, cand)
            if not long:
                continue
            lw = long.split()
            while lw and lw[0].lower() in LONG_FORM_STOP:
                lw.pop(0)
            long = " ".join(lw)
            if not long or len(long) <= len(short) or long.lower() == short.lower():
                continue
            amap.setdefault(short, long)
            if short.endswith("s") and short[:-1] != short[:-1].lower():
                amap.setdefault(short[:-1], long)
            else:
                amap.setdefault(short + "s", long)
    return amap


def apply_acronyms(text: str, amap: dict[str, str]) -> str:
    if not amap or not text:
        return text
    alt = "|".join(re.escape(k) for k in sorted(amap, key=len, reverse=True))
    text = re.sub(r"\s*\(\s*(?:" + alt + r")\s*\)", "", text)
    return re.sub(r"(?<![\w\-])(" + alt + r")(?![\w\-])", lambda m: amap[m.group(1)], text)


def preprocess(df: pd.DataFrame) -> list[dict]:
    docs = []
    for title, abstract, keywords in zip(df[COL_TITLE], df[COL_ABSTRACT], df[COL_KEYWORDS]):
        title = HTML_RE.sub("", title)
        abstract = clean_abstract(abstract)
        kws = [HTML_RE.sub("", k).strip() for k in keywords.split(";") if k.strip()]
        amap = find_acronyms([title, abstract, *kws])
        title = STAT_PAREN_RE.sub("", apply_acronyms(title, amap))
        abstract = STAT_PAREN_RE.sub("", apply_acronyms(abstract, amap))
        kws = [re.sub(r"\s+\([^()]*\)", "", apply_acronyms(k, amap)).strip(" .,") for k in kws]
        docs.append({
            "title": " ".join(title.split()),
            "abstract": " ".join(abstract.split()),
            "keywords": [k for k in kws if k],
        })
    return docs


# ─────────────────────────────────────────────────────────────────────────────
# 3. 품사 태깅 (spaCy) → 단위(unit) 열로 변환
#    unit = (표면형, 분류)  분류: N=명사  A=형용사  O=기타 단어  B=구두점(경계)
# ─────────────────────────────────────────────────────────────────────────────
HYPHENS = {"-", "‐", "‑", "–"}
POSSESSIVE_TOKENS = {"'s", "’s", "'", "’"}


def _token_class(tok) -> str:
    if tok.is_space or tok.is_punct or tok.pos_ in ("PUNCT", "SYM", "SPACE", "X") and not tok.is_alpha:
        return "B"
    if tok.pos_ in ("NOUN", "PROPN"):
        return "N"
    if tok.pos_ == "ADJ":
        return "A"
    return "O"


def doc_to_units(doc) -> list[tuple[str, str]]:
    toks = list(doc)
    units: list[list[str]] = []
    i = 0
    while i < len(toks):
        t = toks[i]
        if t.text in POSSESSIVE_TOKENS:
            i += 1
            continue
        if (t.text in HYPHENS and units and units[-1][1] != "B" and i + 1 < len(toks)
                and not toks[i - 1].whitespace_ and not t.whitespace_ and re.search(r"\w", toks[i + 1].text)):
            nxt_cls = _token_class(toks[i + 1])
            prev_cls = units[-1][1]
            units[-1][0] += "-" + toks[i + 1].text
            units[-1][1] = nxt_cls if nxt_cls in "NA" else (prev_cls if prev_cls in "NA" else "O")
            i += 2
            continue
        units.append([t.text, _token_class(t)])
        i += 1
    return [(s, c) for s, c in units]


def tag_documents(docs: list[dict], processes: int, use_cache: bool, progress=None) -> None:
    import spacy

    try:
        nlp = spacy.load(SPACY_MODEL, exclude=["parser", "ner", "lemmatizer", "senter"])
    except OSError:
        raise SystemExit(f"spaCy 모델이 없습니다. 먼저 실행하세요:  python -m spacy download {SPACY_MODEL}")

    h = hashlib.sha1(f"{CACHE_VERSION}|{SPACY_MODEL}|{nlp.meta.get('version')}".encode())
    for d in docs:
        h.update(d["title"].encode("utf-8", "replace") + b"\x00" + d["abstract"].encode("utf-8", "replace") + b"\x01")
    cache_file = CACHE_DIR / f"tags_{h.hexdigest()[:20]}.pkl"

    if use_cache and cache_file.exists():
        log("[3/5] 품사 태깅: 캐시 사용")
        cached = pickle.loads(cache_file.read_bytes())
    else:
        texts = [t for d in docs for t in (d["title"], d["abstract"])]
        log(f"[3/5] 품사 태깅 중 (spaCy, {len(texts):,}개 텍스트) — 수 분 걸릴 수 있습니다")
        stream = nlp.pipe(texts, batch_size=64, n_process=processes)
        if progress is None:
            from tqdm import tqdm

            stream = tqdm(stream, total=len(texts), unit="text", ncols=80)
        cached = []
        for n, doc in enumerate(stream, 1):
            cached.append(doc_to_units(doc))
            if progress and (n % 200 == 0 or n == len(texts)):
                progress(n, len(texts))
        CACHE_DIR.mkdir(exist_ok=True)
        cache_file.write_bytes(pickle.dumps(cached, protocol=pickle.HIGHEST_PROTOCOL))

    for i, d in enumerate(docs):
        d["title_units"], d["abstract_units"] = cached[2 * i], cached[2 * i + 1]


# ─────────────────────────────────────────────────────────────────────────────
# 4. 정규화 (단수형 · 소문자 · 하이픈/소유격 통일)
# ─────────────────────────────────────────────────────────────────────────────
POSSESSIVE_RE = re.compile(r"(?<=\w)['’]s\b")
WORD_RE = re.compile(r"\d+(?:\.\d+)+|[^\W_]+")


def split_words(s: str) -> list[str]:
    return WORD_RE.findall(POSSESSIVE_RE.sub("", s))


class Normalizer:
    """단어를 단수형 소문자 키로 변환. 사전에 없는 단어는 코퍼스에 단수형이 실제 등장할 때만 변환."""

    def __init__(self, vocab: Counter):
        from lemminflect import getLemma

        self.getLemma, self.vocab, self.cache = getLemma, vocab, {}

    def word(self, w: str) -> str:
        r = self.cache.get(w)
        if r is None:
            r = self.cache[w] = self._word(w)
        return r

    def _word(self, w: str) -> str:
        stem = w[:-1]
        if len(w) >= 3 and w[-1] == "s" and stem == stem.upper() and sum(c.isupper() for c in stem) >= 2:
            return stem.lower()  # LLMs → llm, MSCs → msc
        lw = w.lower()
        if (len(w) >= 2 and w.isupper()) or not lw.isalpha() or len(lw) <= 3:
            return lw  # AIDS, SARS 같은 약어는 그대로
        cands = self.getLemma(lw, upos="NOUN", lemmatize_oov=False)
        if cands:
            return lw if lw in cands else max(cands, key=lambda c: self.vocab.get(c, 0))
        if lw.endswith("s") and not lw.endswith(("ss", "us", "is")):
            oov = self.getLemma(lw, upos="NOUN", lemmatize_oov=True)
            cand = oov[0].lower() if oov else lw
            if cand != lw and self.vocab.get(cand, 0) > 0:
                return cand
        return lw

    def key(self, s: str) -> str:
        return " ".join(self.word(w) for w in split_words(s))


def raw_key(s: str) -> str:
    return " ".join(w.lower() for w in split_words(s))


# ─────────────────────────────────────────────────────────────────────────────
# 5. 설정 파일
# ─────────────────────────────────────────────────────────────────────────────
def read_config_lines(name: str) -> list[str]:
    path = CONFIG_DIR / name
    if not path.exists():
        return []
    lines = []
    for line in path.read_text(encoding="utf-8-sig").splitlines():
        line = re.split(r"\s#|^#", line, maxsplit=1)[0].strip()
        if line:
            lines.append(line)
    return lines


class Config:
    def __init__(self, norm: Normalizer):
        from spacy.lang.en.stop_words import STOP_WORDS

        k = norm.key
        self.function_words = {w.lower() for w in STOP_WORDS}
        self.generic = {k(x) for x in read_config_lines("stopwords.txt")}
        self.modifiers = {k(x) for x in read_config_lines("modifiers.txt")}
        self.trailing = {k(x) for x in read_config_lines("trailing_generic.txt")}
        self.exclude = {k(x) for x in read_config_lines("exclude_phrases.txt")}
        self.custom = {k(x) for x in read_config_lines("custom_phrases.txt")}
        self.synonyms = {}
        for line in read_config_lines("synonyms.txt"):
            if "=>" in line:
                left, right = line.split("=>", 1)
                for variant in left.split("|"):
                    if k(variant) and k(right):
                        self.synonyms[k(variant)] = k(right)
        self.weak = self.generic | self.modifiers | self.function_words

    def accept(self, key: str) -> bool:
        if not key or key in self.exclude:
            return False
        words = key.split()
        if len(words) == 1:
            w = words[0]
            return (w not in self.weak and len(w) >= 2 and any(c.isalpha() for c in w))
        return not all(w in self.weak or not any(c.isalpha() for c in w) for w in words)


# ─────────────────────────────────────────────────────────────────────────────
# 6. 구문 사전 구축 & 집계
# ─────────────────────────────────────────────────────────────────────────────
def unit_runs(keys, classes, cfg: Config):
    """형용사·명사가 연속된 구간(명사구 후보)의 인덱스 목록."""
    run = []
    for i, (k, c) in enumerate(zip(keys, classes)):
        if c in "NA" and k and k not in cfg.function_words and any(ch.isalpha() for ch in k):
            run.append(i)
        else:
            if len(run) >= 2:
                yield run
            run = []
    if len(run) >= 2:
        yield run


def build_phrase_dict(docs, cfg: Config, min_papers: int, max_units: int):
    ngram_papers = Counter()
    for d in docs:
        seen = set()
        for sec in ("title", "abstract"):
            keys, classes = d[f"{sec}_keys"], d[f"{sec}_classes"]
            for run in unit_runs(keys, classes, cfg):
                for a in range(len(run)):
                    if keys[run[a]] in cfg.modifiers:
                        continue
                    for b in range(a + 2, min(a + max_units, len(run)) + 1):
                        last = run[b - 1]
                        if classes[last] != "N" or keys[last] in cfg.trailing or keys[last] in cfg.modifiers:
                            continue
                        seen.add(" ".join(keys[j] for j in run[a:b]))
        ngram_papers.update(seen)

    auto = {p for p, n in ngram_papers.items() if n >= min_papers and cfg.accept(p)}
    from_keywords = {k for d in docs for k in d["keyword_keys"] if 2 <= len(k.split()) <= 8}
    phrases = (auto | from_keywords | cfg.custom) - cfg.exclude
    log(f"      구문 사전: 자동추출 {len(auto):,} + 저자키워드 {len(from_keywords):,} + 사용자 {len(cfg.custom):,}"
          f" → 총 {len(phrases):,}개")
    return phrases


def match_units(keys, classes, surfaces, phrases, first_words, max_words):
    """가장 긴 구문 우선 매칭. 사전에 없는 명사는 단어로 반환."""
    i, n = 0, len(keys)
    while i < n:
        if classes[i] == "B" or not keys[i]:
            i += 1
            continue
        best = None
        if keys[i].split(" ", 1)[0] in first_words:
            joined, words = "", 0
            for j in range(i, n):
                if classes[j] == "B" or not keys[j]:
                    break
                joined = keys[j] if j == i else f"{joined} {keys[j]}"
                words += keys[j].count(" ") + 1
                if words > max_words:
                    break
                if (j > i or " " in joined) and joined in phrases:
                    best = (j, joined)
        if best:
            j, key = best
            yield key, " ".join(surfaces[i:j + 1])
            i = j + 1
            continue
        if classes[i] == "N":
            yield keys[i], surfaces[i]
        i += 1


def count_phrases(docs, cfg: Config, phrases, sections, paper_matches=None):
    first_words = {p.split(" ", 1)[0] for p in phrases}
    max_words = max((len(p.split()) for p in phrases), default=1)
    stats = defaultdict(lambda: {"title": 0, "abstract": 0, "keywords": 0, "papers": 0, "kw_papers": 0})
    forms = defaultdict(Counter)

    for paper_id, d in enumerate(docs):
        per_paper = defaultdict(Counter)
        doc_seen, doc_kw_seen = set(), set()
        items = []
        for sec in ("title", "abstract"):
            if sec in sections:
                items += [(sec, k, s) for k, s in match_units(
                    d[f"{sec}_keys"], d[f"{sec}_classes"], d[f"{sec}_surfaces"], phrases, first_words, max_words)]
        if "keywords" in sections:
            items += [("keywords", k, s) for k, s in zip(d["keyword_keys"], d["keywords"])]
        for sec, key, surface in items:
            key = cfg.synonyms.get(key, key)
            if not cfg.accept(key):
                continue
            stats[key][sec] += 1
            per_paper[key][sec] += 1
            forms[key][surface] += 1
            doc_seen.add(key)
            if sec == "keywords":
                doc_kw_seen.add(key)
        if paper_matches is not None:
            for key, counts in per_paper.items():
                paper_matches.setdefault(key, []).append({"id": paper_id, **dict(counts)})
        for key in doc_seen:
            stats[key]["papers"] += 1
        for key in doc_kw_seen:
            stats[key]["kw_papers"] += 1
    return stats, forms


def display_form(key: str, forms: Counter) -> tuple[str, str]:
    by_lower = defaultdict(Counter)
    for f, c in forms.items():
        by_lower[f.lower()][f] += c
    singular = [l for l in by_lower if raw_key(l) == key]
    best = max(singular or by_lower, key=lambda l: sum(by_lower[l].values()))
    variants = by_lower[best]
    shown = best if variants.get(best, 0) >= 0.25 * sum(variants.values()) else variants.most_common(1)[0][0]
    ordered = sorted(by_lower, key=lambda l: -sum(by_lower[l].values()))
    merged = "; ".join(by_lower[l].most_common(1)[0][0] for l in ordered[:8])
    return shown, merged


# ─────────────────────────────────────────────────────────────────────────────
# 7. 출력 (엑셀 + 워드클라우드)
# ─────────────────────────────────────────────────────────────────────────────
# 다채로운 색 (흰 배경에서 작은 글자도 읽히는 명도)
PALETTE = ["#1f5fbf", "#d1343f", "#e07b00", "#1e9e5a", "#7b3fb8", "#0f8a9d",
           "#c2185b", "#5d7f00", "#b45f06", "#34495e", "#e04e9a", "#2e86de"]
# 글자 각도 후보 (0이 여러 번 들어 있으면 가로가 그만큼 자주 나옴)
ANGLE_SETS = {
    "mixed": [0, 0, 0, 0, 30, -30, 45, -45, 60, -60, 90],
    "right": [0, 0, 90],
    "horizontal": [0],
}


def default_font() -> str | None:
    for p in ("C:/Windows/Fonts/seguisb.ttf", "C:/Windows/Fonts/arialbd.ttf", "C:/Windows/Fonts/arial.ttf",
              "/Library/Fonts/Arial.ttf", "/System/Library/Fonts/Supplemental/Arial.ttf",
              "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf", "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"):
        if Path(p).exists():
            return p
    return None


def _font_file(args) -> str:
    path = args.font or default_font()
    if path:
        return path
    from wordcloud.wordcloud import FONT_PATH  # wordcloud 패키지에 들어 있는 기본 글꼴

    return FONT_PATH


def _color_picker(args, rng):
    if args.colormap:
        import matplotlib

        cmap = matplotlib.colormaps[args.colormap]
        return lambda: tuple(int(255 * c) for c in cmap(rng.uniform(0.05, 0.95))[:3])
    palette, order = PALETTE, []

    def pick():  # 팔레트를 섞어 한 바퀴씩 돌려 색이 고르게 나오도록
        if not order:
            order.extend(rng.sample(palette, len(palette)))
        return order.pop()

    return pick


def _word_sprite(text: str, font, angle: int, cell: int):
    """글자를 angle만큼 돌린 알파 이미지와, 겹침 검사용 격자 마스크(주변 1칸 여백 포함)."""
    import numpy as np
    from PIL import Image, ImageDraw

    left, top, right, bottom = font.getbbox(text)
    pad = 2 * cell
    sprite = Image.new("L", (right - left + 2 * pad, bottom - top + 2 * pad), 0)
    ImageDraw.Draw(sprite).text((pad - left, pad - top), text, font=font, fill=255)
    if angle:
        sprite = sprite.rotate(angle, resample=Image.Resampling.BICUBIC, expand=True)
    ink = np.asarray(sprite) > 32
    gh, gw = -(-ink.shape[0] // cell), -(-ink.shape[1] // cell)
    padded = np.zeros((gh * cell, gw * cell), dtype=bool)
    padded[: ink.shape[0], : ink.shape[1]] = ink
    grid = padded.reshape(gh, cell, gw, cell).any(axis=(1, 3))
    mask = grid.copy()
    mask[1:] |= grid[:-1]
    mask[:-1] |= grid[1:]
    mask[:, 1:] |= grid[:, :-1]
    mask[:, :-1] |= grid[:, 1:]
    return sprite, mask


def make_wordcloud(freqs: dict[str, float], path: Path, args):
    """워드클라우드 이미지 생성.
    wordcloud 패키지는 0°/90°만 지원하므로, 여러 각도 배치를 위해 직접 배치한다.
    글자 크기는 wordcloud 패키지와 같은 방식(relative_scaling), 빈 자리는 가운데에 가까운 곳부터 채운다."""
    import random

    import numpy as np
    from PIL import Image, ImageFont

    items = sorted(freqs.items(), key=lambda kv: -kv[1])
    if not items:
        return
    rng, noise = random.Random(42), np.random.default_rng(42)
    cell, min_size, rel = 4, 9, 0.45
    width, height = args.width, args.height
    gh, gw = height // cell, width // cell
    occupied = np.zeros((gh, gw))
    occupied_f = np.fft.rfft2(occupied)
    font_file, fonts = _font_file(args), {}

    def font(size):
        if size not in fonts:
            fonts[size] = ImageFont.truetype(font_file, size)
        return fonts[size]

    def find_spot(mask):
        """겹치지 않는 위치 중 가운데에 가장 가까운 곳 (FFT 상관으로 모든 위치를 한 번에 검사)."""
        h, w = mask.shape
        if h > gh or w > gw:
            return None
        kernel = np.zeros((gh, gw))
        kernel[:h, :w] = mask
        overlap = np.fft.irfft2(occupied_f * np.conj(np.fft.rfft2(kernel)), s=(gh, gw))
        free = overlap[: gh - h + 1, : gw - w + 1] < 0.5
        if not free.any():
            return None
        ys = (np.arange(gh - h + 1) + h / 2 - gh / 2) / gh
        xs = (np.arange(gw - w + 1) + w / 2 - gw / 2) / gw
        score = ys[:, None] ** 2 + xs[None, :] ** 2 + noise.uniform(0, 0.004, free.shape)
        score[~free] = np.inf
        y, x = np.unravel_index(np.argmin(score), score.shape)
        return int(y), int(x)

    size = int(height * 0.5)  # 1위 항목이 가로로 화면에 들어가는 크기에서 시작
    while size > min_size:
        left, top, right, bottom = font(size).getbbox(items[0][0])
        if right - left <= width * 0.6 and bottom - top <= height * 0.28:
            break
        size = int(size * 0.92)

    angles, pick_color = ANGLE_SETS[args.angles], _color_picker(args, rng)
    canvas = Image.new("RGB", (width, height), args.background)
    hitmap = Image.new("RGB", (width, height), (0, 0, 0))
    placements = []
    last_freq = items[0][1]
    for rank, (word, freq) in enumerate(items):
        if rank:
            size = int(round((rel * freq / last_freq + (1 - rel)) * size))
        last_freq = freq
        angle = 0 if rank < 3 else rng.choice(angles)  # 상위 3개는 가로로
        placed = False
        while size >= min_size and not placed:
            for a in dict.fromkeys([angle, 0]):  # 자리가 없으면 가로로 한 번 더 시도
                sprite, mask = _word_sprite(word, font(size), a, cell)
                spot = find_spot(mask)
                if spot:
                    y, x = spot
                    canvas.paste(pick_color(), (x * cell, y * cell), sprite)
                    # 원본과 동일한 이미지 배치에 클릭용 픽셀 ID만 추가한다.
                    ident = len(placements) + 1
                    hitmap.paste((ident & 255, (ident >> 8) & 255, (ident >> 16) & 255),
                                 (x * cell, y * cell), sprite.point(lambda v: 255 if v > 32 else 0))
                    placements.append({"word": word, "x": x * cell, "y": y * cell,
                                       "w": sprite.width, "h": sprite.height, "angle": a})
                    occupied[y:y + mask.shape[0], x:x + mask.shape[1]] += mask
                    occupied_f = np.fft.rfft2(occupied)
                    placed = True
                    break
            else:
                size = int(size * 0.9)
        if not placed:
            break  # 가장 작은 글자도 들어갈 자리가 없음
    canvas.save(str(path))
    from paper_explorer import png_data
    return {"png": png_data(canvas), "hit": png_data(hitmap),
            "placements": placements, "width": width, "height": height}


def autosize(ws, df: pd.DataFrame) -> None:
    from openpyxl.utils import get_column_letter

    for i, col in enumerate(df.columns, start=1):
        sample = df[col].astype(str).head(300)
        width = max([len(str(col)) * 2] + [len(v) for v in sample]) + 2
        ws.column_dimensions[get_column_letter(i)].width = min(max(width, 8), 70)
    ws.freeze_panes = "A2"


def build_table(stats, forms, cfg: Config, args) -> tuple[pd.DataFrame, pd.DataFrame]:
    """키워드성 필터: 등장 논문 중 저자키워드로 쓴 논문 비율이 기준 이상인 항목만 채택.
    ('risk', 'valuable insight' 처럼 초록에 흔하지만 연구 주제어로는 쓰이지 않는 일반어 제거)"""
    rows = []
    for key, s in stats.items():
        if s["papers"] < args.min_papers:
            continue
        ratio = s["kw_papers"] / s["papers"]
        passed = key in cfg.custom or (s["kw_papers"] >= args.min_keyword_papers and ratio >= args.min_keyword_ratio)
        shown, merged = display_form(key, forms[key])
        rows.append({
            "키프레이즈": shown, "단어수": len(key.split()),
            "총빈도": s["title"] + s["abstract"] + s["keywords"], "논문수": s["papers"],
            "제목": s["title"], "초록": s["abstract"], "저자키워드": s["keywords"],
            "저자키워드 사용률": round(ratio, 3), "병합된 표기": merged, "정규화 키": key, "_pass": passed,
        })
    sort_col, other = ("총빈도", "논문수") if args.weight == "count" else ("논문수", "총빈도")
    table = pd.DataFrame(rows).sort_values([sort_col, other], ascending=False, ignore_index=True)
    kept = table[table["_pass"]].drop(columns="_pass").reset_index(drop=True)
    dropped = table[~table["_pass"]].drop(columns="_pass").reset_index(drop=True)
    return kept, dropped


def write_outputs(table: pd.DataFrame, dropped: pd.DataFrame, settings: dict, args, out_dir: Path,
                  label: str, paper_matches=None, records=None) -> list[Path]:
    out_dir.mkdir(parents=True, exist_ok=True)
    weight_col = "총빈도" if args.weight == "count" else "논문수"
    xlsx = out_dir / f"{label}_키프레이즈_빈도.xlsx"
    sheets = {
        "전체": table,
        "구문(2단어 이상)": table[table["단어수"] >= 2],
        "단어(1단어)": table[table["단어수"] == 1],
        "제외됨(일반어 필터)": dropped.head(3000),
    }
    with pd.ExcelWriter(xlsx, engine="openpyxl") as xw:
        for name, t in sheets.items():
            t = t.reset_index(drop=True)
            t.insert(0, "순위", range(1, len(t) + 1))
            t.to_excel(xw, sheet_name=name, index=False)
            autosize(xw.sheets[name], t)
        s = pd.DataFrame(list(settings.items()), columns=["항목", "값"])
        s.to_excel(xw, sheet_name="설정", index=False)
        autosize(xw.sheets["설정"], s)

    outputs = [xlsx]
    clouds = []
    for suffix, t in (("전체", sheets["전체"]), ("구문", sheets["구문(2단어 이상)"])):
        top = t.head(args.top)
        freqs = defaultdict(float)
        for name, w in zip(top["키프레이즈"], top[weight_col]):
            freqs[name] += w
        if freqs:
            png = out_dir / f"{label}_워드클라우드_{suffix}.png"
            cloud = make_wordcloud(dict(freqs), png, args)
            cloud["label"] = "전체 (단어 + 구문)" if suffix == "전체" else "구문만 (2단어 이상)"
            clouds.append(cloud)
            outputs.append(png)
    if clouds and paper_matches is not None and records is not None:
        from paper_explorer import write_report
        report = out_dir / f"{label}_논문탐색.html"
        write_report(report, clouds, table, paper_matches, records, label)
        outputs.append(report)
    return outputs


# ─────────────────────────────────────────────────────────────────────────────
# main
# ─────────────────────────────────────────────────────────────────────────────
def parse_args(argv=None):
    ap = argparse.ArgumentParser(description="Scopus 논문 키프레이즈 빈도 집계 + 워드클라우드")
    ap.add_argument("inputs", nargs="+", help="Scopus 내보내기 파일 (.xlsx / .csv). 여러 개 지정 가능")
    ap.add_argument("--sheet", help="읽을 시트 이름 (기본: Title/Abstract/Author Keywords 열이 있는 첫 시트)")
    ap.add_argument("--year", type=int, nargs="*", help="특정 발행연도만 집계 (예: --year 2026)")
    ap.add_argument("--label", help="결과 파일 이름 앞부분 (기본: 입력 파일 이름)")
    ap.add_argument("--out", default=str(BASE_DIR / "output"), help="결과 저장 폴더")
    ap.add_argument("--sections", nargs="+", choices=["title", "abstract", "keywords"],
                    default=["title", "abstract", "keywords"], help="집계할 섹션")
    ap.add_argument("--weight", choices=["count", "papers"], default="count",
                    help="워드클라우드 크기 기준: count=총 노출빈도(기본), papers=등장 논문 수")
    ap.add_argument("--top", type=int, default=150, help="워드클라우드에 표시할 키프레이즈 수")
    ap.add_argument("--min-phrase-papers", type=int, default=3,
                    help="제목/초록에서 자동 구문으로 인정할 최소 등장 논문 수")
    ap.add_argument("--max-phrase-units", type=int, default=4, help="자동 구문의 최대 단어(토큰) 수")
    ap.add_argument("--min-papers", type=int, default=2, help="엑셀 표에 넣을 최소 등장 논문 수")
    ap.add_argument("--min-keyword-ratio", type=float, default=0.05,
                    help="키워드성 필터: 등장 논문 중 저자키워드로 사용한 논문 비율 하한 (0이면 필터 끔)")
    ap.add_argument("--min-keyword-papers", type=int, default=2,
                    help="키워드성 필터: 저자키워드로 사용한 최소 논문 수 (0이면 필터 끔)")
    ap.add_argument("--processes", type=int, default=1, help="spaCy 병렬 프로세스 수")
    ap.add_argument("--no-cache", action="store_true", help="품사 태깅 캐시를 사용하지 않음")
    ap.add_argument("--width", type=int, default=2000)
    ap.add_argument("--height", type=int, default=1200)
    ap.add_argument("--background", default="white")
    ap.add_argument("--font", help="워드클라우드 글꼴(.ttf) 경로")
    ap.add_argument("--angles", choices=list(ANGLE_SETS), default="mixed",
                    help="워드클라우드 글자 각도: mixed=다양한 각도(기본), right=가로+세로, horizontal=가로만")
    ap.add_argument("--colormap", help="matplotlib 컬러맵 이름 (예: viridis). 지정하면 기본 12색 팔레트 대신 사용")
    return ap.parse_args(argv)


def run_pipeline(args, progress=None) -> dict:
    """전체 분석 실행. CLI(main)와 웹 화면(app.py)이 함께 사용."""
    label = args.label or Path(args.inputs[0]).stem
    started = dt.datetime.now()

    df, removed = load_records(args.inputs, args.sheet, args.year)
    if df.empty:
        raise SystemExit("집계할 논문이 없습니다.")
    log(f"      논문 {len(df):,}편 (중복·빈 행 {removed:,}건 제외)")

    log("[2/5] 전처리 중 (저작권 문구·구조화 초록 제목·약어 처리)")
    docs = preprocess(df)
    tag_documents(docs, args.processes, not args.no_cache, progress)

    log("[4/5] 키프레이즈 추출·집계 중")
    vocab = Counter()
    for d in docs:
        for s, _ in d["title_units"] + d["abstract_units"]:
            vocab.update(w.lower() for w in split_words(s))
        for k in d["keywords"]:
            vocab.update(w.lower() for w in split_words(k))
    norm = Normalizer(vocab)
    cfg = Config(norm)
    for d in docs:
        for sec in ("title", "abstract"):
            units = d.pop(f"{sec}_units")
            d[f"{sec}_surfaces"] = [s for s, _ in units]
            d[f"{sec}_classes"] = [c for _, c in units]
            d[f"{sec}_keys"] = [norm.key(s) if c != "B" else "" for s, c in units]
        d["keyword_keys"] = [norm.key(k) for k in d["keywords"]]

    phrases = build_phrase_dict(docs, cfg, args.min_phrase_papers, args.max_phrase_units)
    paper_matches = {}
    stats, forms = count_phrases(docs, cfg, phrases, set(args.sections), paper_matches)
    table, dropped = build_table(stats, forms, cfg, args)
    log(f"      키프레이즈 {len(table):,}개 채택, 일반어 필터로 {len(dropped):,}개 제외")

    log("[5/5] 결과 저장 중")
    settings = {
        "입력 파일": " | ".join(Path(p).name for p in args.inputs),
        "시트": args.sheet or "(자동 선택)",
        "연도 필터": ", ".join(map(str, args.year)) if args.year else "(없음)",
        "집계 논문 수": len(df),
        "제외된 중복/빈 행": removed,
        "초록 없는 논문": sum(1 for d in docs if not d["abstract"]),
        "저자키워드 없는 논문": sum(1 for d in docs if not d["keywords"]),
        "집계 섹션": ", ".join(SECTION_LABELS[s] for s in args.sections),
        "워드클라우드 기준": "총빈도(노출 횟수)" if args.weight == "count" else "논문수",
        "워드클라우드 표시 개수": args.top,
        "워드클라우드 각도": args.angles,
        "워드클라우드 색상": args.colormap or "다채로운 색 (12색)",
        "자동 구문 최소 논문 수": args.min_phrase_papers,
        "자동 구문 최대 단어 수": args.max_phrase_units,
        "표 포함 최소 논문 수": args.min_papers,
        "키워드성 필터: 저자키워드 사용률 하한": args.min_keyword_ratio,
        "키워드성 필터: 저자키워드 사용 최소 논문 수": args.min_keyword_papers,
        "구문 사전 크기": len(phrases),
        "실행 일시": started.strftime("%Y-%m-%d %H:%M"),
    }
    outputs = write_outputs(table, dropped, settings, args, Path(args.out) / label, label, paper_matches, df)
    return {
        "label": label, "table": table, "dropped": dropped, "settings": settings, "outputs": outputs,
        "weight_col": "총빈도" if args.weight == "count" else "논문수",
        "seconds": (dt.datetime.now() - started).seconds,
    }


def main():
    res = run_pipeline(parse_args())
    table, weight_col = res["table"], res["weight_col"]
    print(f"\n상위 30개 키프레이즈 ({weight_col} 기준)")
    for i, r in table.head(30).iterrows():
        print(f"  {i + 1:>3}. {r['키프레이즈']:<40} 총빈도 {r['총빈도']:>6,}  논문 {r['논문수']:>5,}")
    print("\n저장된 파일:")
    for p in res["outputs"]:
        print("  -", p)
    print(f"완료 ({res['seconds']}초)")


if __name__ == "__main__":
    main()
